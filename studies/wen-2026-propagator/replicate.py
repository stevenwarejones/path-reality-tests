#!/usr/bin/env python3
"""Replot deposited summaries and compute descriptive residuals, without fitting.

No full propagator tensor, repeat-level data, or significance is reconstructed.
All input bytes are verified before output; originals are opened read-only.
"""
import argparse
import csv
import io
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl

from fetch_data import ROOT, verify
from synthetic_checks import compare_snapshot, smape


def block(rows, columns, count):
    """Read a named contiguous numeric block; reject holes and trailing values."""
    indexes = [ord(c) - ord("A") for c in columns]
    selected = [[row[i] for i in indexes] for row in rows[1:count + 1]]
    if len(selected) != count or any(
            isinstance(v, bool) or not isinstance(v, (float, int))
            for row in selected for v in row):
        raise ValueError(f"Missing/nonnumeric value in {columns}, rows 2:{count + 1}")
    result = np.asarray(selected, dtype=float)
    if not np.isfinite(result).all():
        raise ValueError("Nonfinite numeric value")
    if any(row[i] is not None for row in rows[count + 1:] for i in indexes):
        raise ValueError(f"Unexpected trailing values in {columns}")
    return result


def align_theory(x, y, target, method="linear"):
    x, y, target = map(lambda a: np.asarray(a, dtype=float), (x, y, target))
    if (x.ndim != 1 or y.shape != x.shape or len(x) < 2
            or not all(np.isfinite(a).all() for a in (x, y, target))
            or np.any(np.diff(x) <= 0)):
        raise ValueError("Theory coordinates must be finite, unique and increasing")
    if np.any(target < x[0]) or np.any(target > x[-1]):
        raise ValueError("Extrapolation is not permitted")
    if method == "linear":
        return np.interp(target, x, y)
    if method == "nearest":
        return y[np.argmin(np.abs(x[:, None] - target), axis=0)]
    if method == "local_cubic":
        if len(x) < 4:
            raise ValueError("Local cubic interpolation requires four points")
        values = []
        for z in target:
            start = min(max(int(np.searchsorted(x, z)) - 2, 0), len(x) - 4)
            values.append(np.polynomial.polynomial.polyfit(
                x[start:start + 4] - z, y[start:start + 4], 3)[0])
        return np.array(values)
    raise ValueError(f"Unknown alignment: {method}")


def metrics(a, b, percentage=True):
    a, b = np.asarray(a, dtype=float), np.asarray(b, dtype=float)
    if a.shape != b.shape or a.size == 0 or not np.isfinite([a, b]).all():
        raise ValueError("Comparisons need matching nonempty finite arrays")
    d = a - b
    result = {"n": int(a.size), "mean_signed_difference": float(d.mean()),
            "mae": float(np.abs(d).mean()), "rmse": float(np.sqrt(np.mean(d ** 2))),
            "max_absolute_difference": float(np.abs(d).max())}
    if percentage:
        result["smape_percent"] = 100 * smape(a, b)
    return result


def wrap_phase_pi(delta):
    """Shortest signed angular difference in units of pi: [-1, 1)."""
    return (np.asarray(delta) + 1) % 2 - 1


def source_records(entries, names):
    return [{k: entries[name][k] for k in ("name", "size", "sha256")}
            for name in names]


def inspect_inputs(directory):
    manifest = json.loads((ROOT / "manifest.json").read_text())
    entries = {e["name"]: e for e in manifest["files"]}
    for entry in entries.values():
        verify(directory / entry["name"], entry)
    mapping = json.loads((ROOT / "column-map.json").read_text())
    books, audit = {}, []
    for name, sheets in mapping.items():
        book = openpyxl.load_workbook(directory / name, read_only=True, data_only=False)
        try:
            if book.sheetnames != list(sheets):
                raise ValueError(f"Unexpected sheet set: {name}")
            books[name] = {}
            for ws in book:
                spec = sheets[ws.title]
                rows = list(ws.values)
                if [ws.max_row, ws.max_column] != spec["shape"] or ws.sheet_state != "visible":
                    raise ValueError(f"Unexpected sheet dimensions/state: {name}/{ws.title}")
                expected_header = spec.get("first_row")
                if "first_row_axis" in spec:
                    axis = spec["first_row_axis"]
                    expected_header = [None] + [axis["start"] + i * axis["step"] for i in range(axis["count"])]
                    expected_header += [None] * (ws.max_column - len(expected_header))
                if list(rows[0]) != expected_header:
                    raise ValueError(f"Unexpected headers: {name}/{ws.title}")
                formulas = sum(c.data_type == "f" for row in ws for c in row)
                if formulas:
                    raise ValueError(f"Formula present in value-only deposit: {name}")
                numeric = [v for row in rows[1:] for v in row if isinstance(v, (int, float))]
                if not np.isfinite(numeric).all():
                    raise ValueError(f"Nonfinite data: {name}")
                books[name][ws.title] = rows
                for b in spec.get("blocks", []):
                    block(rows, b["columns"], b["count"])
                audit.append({"file": name, "sha256": entries[name]["sha256"],
                              "sheet": ws.title, "shape": spec["shape"],
                              "state": ws.sheet_state, "formulas": formulas,
                              "numeric_cells_below_header": len(numeric),
                              "nonempty_cells": sum(v is not None for row in rows for v in row),
                              "role": spec["role"]})
        finally:
            book.close()
    if set(mapping) != {n for n in entries if n.endswith(".xlsx")}:
        raise ValueError("Column map does not cover every workbook")
    return entries, books, audit


def numeric_sheet(books, name, columns, count, sheet="Sheet1"):
    return block(books[name][sheet], columns, count)


def analyze(directory):
    entries, books, audit = inspect_inputs(directory)
    theory = numeric_sheet(books, "Fig3.xlsx", "ABC", 850)
    obs = numeric_sheet(books, "Fig3.xlsx", "DEFGHI", 17)
    x, e, q, se_q, c, se_c = obs.T
    if np.any(obs[:, 1:] < 0):
        raise ValueError("Probabilities and SEs must be nonnegative")
    qt = align_theory(theory[:, 0], theory[:, 1], x)
    ct = align_theory(theory[:, 0], theory[:, 2], x)
    comparisons = {"E_vs_Q": metrics(e, q), "Q_vs_deposited_theory": metrics(q, qt),
                   "E_vs_C": metrics(e, c), "C_vs_deposited_theory": metrics(c, ct)}
    sensitivity = {}
    for method in ("linear", "nearest", "local_cubic"):
        aligned = align_theory(theory[:, 0], theory[:, 1], x, method)
        sensitivity[method] = {
            "original_units": metrics(q, aligned),
            "unit_sum_shape_only": metrics(q / q.sum(), aligned / aligned.sum()),
            "theory_scale_to_match_Q_sum": float(q.sum() / aligned.sum())}
    endpoint_shape = {"E_vs_Q_unit_sum": metrics(e / e.sum(), q / q.sum()),
                      "E_vs_C_unit_sum": metrics(e / e.sum(), c / c.sum()),
                      "positions_E_closer_to_Q_than_C": int(np.sum(np.abs(e - q) < np.abs(e - c)))}
    definitions = {"ordinary_MAPE_theory_denominator_percent": float(100 * np.mean(np.abs(q - qt) / qt)),
                   "ordinary_MAPE_Q_denominator_percent": float(100 * np.mean(np.abs(q - qt) / q)),
                   "pooled_symmetric_absolute_percent": float(200 * np.sum(np.abs(q - qt)) / np.sum(q + qt))}
    # A conditional bound using the tabulated maximum slope and declared display
    # precision; it is not a bound on unknown analytic curvature or source code.
    curve_value_bound = float(np.max(np.abs(np.diff(theory[:, 1]) / np.diff(theory[:, 0]))) * 5e-6 + .5e-9)
    rounding_metric_bound_pp = float(100 * np.mean(
        4 * (q + .5e-9) / (q + qt - curve_value_bound - .5e-9) ** 2) * curve_value_bound
        + 100 * np.mean(4 * (qt + curve_value_bound)
                        / (q + qt - curve_value_bound - .5e-9) ** 2) * .5e-9)
    # A conditional finite-grid model, not a reconstruction of the apparatus.
    finite_grid = []
    grid = np.arange(-8, 9)
    for dx in (5.73e-6, 6.09e-6):
        kernel = np.exp(1j * np.pi * dx ** 2 / (795e-9 * 15e-3)
                        * (grid[:, None] - grid[None, :]) ** 2)
        probability = np.abs(np.linalg.matrix_power(kernel, 5)[:, 8]) ** 2
        finite_grid.append({"dx_m": dx, "comparison": "unit-sum shapes only",
                            "Q_vs_conditional_grid": metrics(q / q.sum(), probability / probability.sum()),
                            "deposited_theory_vs_conditional_grid": metrics(
                                qt / qt.sum(), probability / probability.sum())})
    fa = numeric_sheet(books, "Fig4A.xlsx", "ABC", 72)
    fb = numeric_sheet(books, "Fig4B.xlsx", "ABC", 100)
    fc = numeric_sheet(books, "Fig4C.xlsx", "ABCD", 100)
    if any(np.any(data[:, 2] < 0) for data in (fa, fb, fc)):
        raise ValueError("SDs must be nonnegative")
    if not np.array_equal(fb[:, 0], fc[:, 0]):
        raise ValueError("Figure 4 action axes do not agree")
    phase_delta = fc[:, 1] - fc[:, 3]
    circular = wrap_phase_pi(phase_delta)
    copies = {
        "Fig4A_vs_FigS4A_varying": float(np.max(np.abs(
            fa - numeric_sheet(books, "FigS4A.xlsx", "ABC", 72)))),
        "Fig4B_vs_FigS4B_varying": float(np.max(np.abs(
            fb - numeric_sheet(books, "FigS4B.xlsx", "ABC", 100)))),
        "Fig4C_vs_FigS4B_varying": float(np.max(np.abs(
            fc - numeric_sheet(books, "FigS4B.xlsx", "ADEF", 100)))),
        "Fig4C_theory_vs_FigS3B_simulation_label": float(np.max(np.abs(
            fc[:, 3] - numeric_sheet(books, "FigS3B.xlsx", "D", 100)[:, 0])))}
    histograms = {}
    for name in ("FigS2A.xlsx", "FigS2C.xlsx"):
        h = np.array(books[name]["Sheet1"][1:], dtype=float)
        counts = h[:, 1:]
        if np.any(counts < 0) or not np.array_equal(counts, np.round(counts)):
            raise ValueError("Histogram counts must be nonnegative integers")
        histograms[name] = {"numeric_rows": len(h),
                            "nonincreasing_coordinate_excel_rows": (np.where(np.diff(h[:, 0]) <= 0)[0] + 3).tolist(),
                            "counts_by_displacement": counts.sum(axis=0).astype(int).tolist(),
                            "total_counts": int(counts.sum())}
        if name == "FigS2C.xlsx":
            # Report both segments; do not silently drop the nonmonotone tail.
            tail = h[99:]
            histograms[name]["rows_101_120_nonzero_count_cells"] = [
                {"cell": f"{openpyxl.utils.get_column_letter(j + 2)}{i + 101}",
                 "count": int(tail[i, j + 1])}
                for i, j in zip(*np.nonzero(tail[:, 1:]))]
            histograms[name]["rows_2_100_counts_total"] = int(counts[:99].sum())
    s1 = numeric_sheet(books, "FigS1C.xlsx", "ABC", 340)
    s1obs = numeric_sheet(books, "FigS1C.xlsx", "DEFGH", 17)
    s1_metrics = {"imaginary_component": metrics(s1obs[:, 1], align_theory(s1[:, 0], s1[:, 1], s1obs[:, 0]), False),
                  "real_component": metrics(s1obs[:, 3], align_theory(s1[:, 0], s1[:, 2], s1obs[:, 0]), False)}
    image_info = {}
    for name in books:
        if not name.startswith("FigS1B-"):
            continue
        rows = books[name]["Sheet1"]
        # The worksheet has formatted blank columns after CX: those are not pixels.
        if any(v is not None for r in rows for v in r[102:]):
            raise ValueError("Unexpected data beyond the image's coordinate header")
        xx, yy = np.array([r[0] for r in rows[1:]]), np.array(rows[0][1:102])
        pixels = np.array([r[1:102] for r in rows[1:]], dtype=float)
        if not np.isfinite(pixels).all():
            raise ValueError("Missing image pixel")
        image_info[name] = {"matrix_shape": list(pixels.shape),
                            "first_column_coordinate_range": [float(xx.min()), float(xx.max())],
                            "first_row_coordinate_range": [float(yy.min()), float(yy.max())],
                            "grayscale_range": [float(pixels.min()), float(pixels.max())]}
    summary = {
        "kind": "Descriptive replication of deposited figure tables; no repeat-level inference",
        "sources": source_records(entries, list(entries)), "workbook_audit": audit,
        "figure3": {"comparisons": comparisons, "alignment_sensitivity": sensitivity,
                    "endpoint_shape_sensitivity": endpoint_shape,
                    "alternative_metric_definitions": definitions,
                    "conditional_plot_rounding_smape_bound_percentage_points": rounding_metric_bound_pp,
                    "deposited_series_means": {"E_17_points": float(e.mean()), "Q_17_points": float(q.mean()),
                                               "C_17_points": float(c.mean()), "Q_theory_850_points": float(theory[:, 1].mean()),
                                               "C_theory_850_points": float(theory[:, 2].mean()),
                                               "Q_theory_interpolated_17_points": float(qt.mean())},
                    "exact_theory_coordinate_matches": int(np.isin(x, theory[:, 0]).sum()),
                    "theory_grid_points": len(theory), "published_Q_theory_percent": 4.45,
                    "Q_theory_difference_from_quote_percentage_points": comparisons["Q_vs_deposited_theory"]["smape_percent"] - 4.45,
                    "relative_difference_from_quoted_error_percent": 100 * (comparisons["Q_vs_deposited_theory"]["smape_percent"] / 4.45 - 1),
                    "coordinate_interpretation_micrometres": {"centers_dx_5_73": [-8 * 5.73, 8 * 5.73],
                                                            "half_bin_edges_dx_5_73": [-8.5 * 5.73, 8.5 * 5.73]},
                    "conditional_finite_grid_sensitivity": finite_grid,
                    "E_uncertainty_available": False},
        "figure4": {"length_rows": len(fa), "length_coordinate_range": [float(fa[0, 0]), float(fa[-1, 0])],
                    "action_rows": len(fb), "action_coordinate_range": [float(fb[0, 0]), float(fb[-1, 0])],
                    "phase_signed_metrics_pi": metrics(fc[:, 1], fc[:, 3], False),
                    "phase_shortest_arc_mae_pi": float(np.mean(np.abs(circular))),
                    "phase_shortest_arc_rmse_pi": float(np.sqrt(np.mean(circular ** 2))),
                    "phase_mean_reported_sd_pi": float(fc[:, 2].mean()),
                    "length_mean_range_au": [float(fa[:, 1].min()), float(fa[:, 1].max())],
                    "action_mean_range_au": [float(fb[:, 1].min()), float(fb[:, 1].max())]},
        "duplicate_summary_max_absolute_differences": copies,
        "histogram_audit": histograms, "sample_propagator_descriptive_residuals": s1_metrics,
        "sample_images": image_info,
        "limitations": ["No complete K[slice,out,in,repeat] tensor is present in the inspected sheets.",
                        "Processed means and SD/SE columns cannot recover path-level fidelities or covariance.",
                        "No fitting, rescaling, dropped endpoints or phase-offset adjustment in primary comparisons.",
                        "Interpolation and unit-sum normalization are declared sensitivities, not the authors' code.",
                        "No p-values: E uncertainties, joint repeats and calibration covariance are unavailable."]}
    return summary, (theory, obs, qt, ct, fa, fb, fc)


def write_json(path, value):
    path.write_text(json.dumps(value, indent=2, allow_nan=False) + "\n")


def csv_text(sources, fields, rows):
    f = io.StringIO(newline="")
    for source in sources:
        f.write(f"# source={source['name']} sha256={source['sha256']}\n")
    writer = csv.writer(f, lineterminator="\n")
    writer.writerow(fields)
    writer.writerows(rows)
    return f.getvalue()


def residual_tables(summary, arrays):
    _, obs, qt, ct, _, _, fc = arrays
    sources = {s["name"]: s for s in summary["sources"]}
    e, q, c = obs[:, 1], obs[:, 2], obs[:, 4]
    delta = fc[:, 1] - fc[:, 3]
    return {"figure3-residuals.csv": csv_text([sources["Fig3.xlsx"]],
                ["x_over_dx", "E_minus_Q_au", "Q_minus_theory_linear_au", "E_minus_C_au", "C_minus_theory_linear_au"],
                zip(obs[:, 0], e - q, q - qt, e - c, c - ct)),
            "figure4-phase-residuals.csv": csv_text([sources["Fig4C.xlsx"]],
                ["action_coordinate_pi_hbar", "mean_phase_minus_theory_pi", "shortest_arc_difference_pi"],
                zip(fc[:, 0], delta, wrap_phase_pi(delta)))}


def save_figure(fig, path, sources):
    credit = "\n".join(f"{s['name']} SHA-256 {s['sha256']}" for s in sources)
    fig.text(.025, .012, credit, fontsize=5.5, family="monospace", va="bottom")
    fig.savefig(path, dpi=180, metadata={"Description": credit,
                "Source": "Dryad 10.5061/dryad.x0k6djj14 version 3 / 450489"})
    plt.close(fig)


def render(out, summary, arrays):
    theory, obs, qt, ct, fa, fb, fc = arrays
    sources = {s["name"]: s for s in summary["sources"]}
    fig, ax = plt.subplots(figsize=(9, 5.8))
    fig.subplots_adjust(bottom=.20, top=.86)
    ax.plot(theory[:, 0], theory[:, 1], color="#1764a0", label="Q theory (deposited)")
    ax.plot(theory[:, 0], theory[:, 2], "--", color="#b86916", label="C theory (deposited)")
    ax.plot(obs[:, 0], obs[:, 1], "^", color="#237a35", label="E: endpoint image")
    ax.errorbar(obs[:, 0], obs[:, 2], yerr=obs[:, 3], fmt="o", color="#1764a0", capsize=2, label="Q: reconstructed ± SE")
    ax.errorbar(obs[:, 0], obs[:, 4], yerr=obs[:, 5], fmt="s", color="#b86916", capsize=2, label="C: incoherent sum ± SE")
    ax.set(xlabel=r"Final coordinate $x_f/\delta x$", ylabel="Probability (a.u.)",
           ylim=(0, 1.5), title="Figure 3 — replot of the deposited summaries")
    ax.legend(fontsize=8, ncol=2, loc="lower center")
    ax.grid(alpha=.18)
    save_figure(fig, out / "figure3-reproduced.png", [sources["Fig3.xlsx"]])
    fig, axes = plt.subplots(1, 3, figsize=(13, 5.2))
    fig.subplots_adjust(bottom=.25, top=.83, wspace=.35)
    for ax, data, title, xlabel in zip(axes, (fa, fb, fc),
            ("A · Length groups", "B · Action groups", "C · Phase versus action"),
            (r"$L/\delta x$", r"Action coordinate ($\pi\hbar$ units)", r"Action coordinate ($\pi\hbar$ units)")):
        ax.fill_between(data[:, 0], data[:, 1] - data[:, 2], data[:, 1] + data[:, 2], color="#5176b8", alpha=.22, label="Deposited ± SD")
        ax.plot(data[:, 0], data[:, 1], ".", color="#294d89", markersize=3, label="Deposited mean")
        ax.set(xlabel=xlabel, title=title)
        ax.grid(alpha=.18)
    axes[0].set_ylabel("Path weight (a.u.)")
    axes[0].set_ylim(0, 1.45)
    axes[1].set_ylabel("Path weight (a.u.)")
    axes[1].set_ylim(0, 1.45)
    axes[2].set_ylabel(r"Phase / $\pi$")
    axes[2].set_ylim(-.1, 2.5)
    axes[2].plot(fc[:, 0], fc[:, 3], "--", color="#9b402f", label="Deposited theory")
    axes[2].legend(fontsize=7)
    fig.suptitle("Figure 4 — deposited means and SDs; coordinates unchanged")
    save_figure(fig, out / "figure4-reproduced.png", [sources[f"Fig4{p}.xlsx"] for p in "ABC"])
    for name, content in residual_tables(summary, arrays).items():
        (out / name).write_text(content)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--directory", type=Path, default=ROOT / "raw" / "dataset")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "results" / "replication")
    parser.add_argument("--check", action="store_true", help="Compare committed numerical results without replacing them")
    args = parser.parse_args()
    if args.output_dir.resolve().is_relative_to(args.directory.resolve()):
        parser.error("Output must be outside the original-data directory")
    summary, arrays = analyze(args.directory)
    if args.check:
        compare_snapshot(summary, json.loads((ROOT / "results" / "replication" / "summary.json").read_text()))
        # Check CSV numerics with the same cross-platform tolerances as JSON.
        for name, content in residual_tables(summary, arrays).items():
            actual = list(csv.reader(content.splitlines()))
            expected = list(csv.reader((ROOT / "results" / "replication" / name).read_text().splitlines()))
            if actual[:2] != expected[:2] or len(actual) != len(expected):
                raise ValueError(f"Residual table source/header/length changed: {name}")
            compare_snapshot([[float(v) for v in row] for row in actual[2:]],
                             [[float(v) for v in row] for row in expected[2:]], name)
        if args.output_dir.resolve() == (ROOT / "results" / "replication").resolve():
            print("Replication summary verified; no tracked outputs rewritten")
            return
    args.output_dir.mkdir(parents=True, exist_ok=True)
    write_json(args.output_dir / "summary.json", summary)
    render(args.output_dir, summary, arrays)
    print("Replicated Figures 3/4; audited all 18 workbooks / 19 sheets")
    print(json.dumps(summary["figure3"]["comparisons"], indent=2))


if __name__ == "__main__":
    main()
