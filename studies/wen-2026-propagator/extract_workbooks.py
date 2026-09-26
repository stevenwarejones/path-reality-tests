#!/usr/bin/env python3
"""Inventory every sheet of verified originals without editing the workbooks.

Stores cell coordinates, formula expressions and cached values separately.
This is extraction only; semantic column mapping still requires inspection.
"""
import argparse
import json
from pathlib import Path
import openpyxl
from fetch_data import ROOT, verify


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--directory", type=Path, default=ROOT / "raw" / "dataset")
    parser.add_argument("--output", type=Path, default=ROOT / "results" / "workbooks")
    args = parser.parse_args()
    entries = [e for e in json.loads((ROOT / "manifest.json").read_text())["files"]
               if e["name"].endswith(".xlsx")]
    # Verify the complete input set before writing any extracts.
    for entry in entries:
        verify(args.directory / entry["name"], entry)
    out = args.output
    out.mkdir(parents=True, exist_ok=True)
    for entry in entries:
        source = args.directory / entry["name"]
        formulas = openpyxl.load_workbook(source, read_only=True, data_only=False)
        values = openpyxl.load_workbook(source, read_only=True, data_only=True)
        report = {"source": entry, "sheets": []}
        for ws in formulas:
            cached = values[ws.title]
            cells = []
            for frow, vrow in zip(ws.iter_rows(), cached.iter_rows(), strict=True):
                for cell, value in zip(frow, vrow, strict=True):
                    if cell.value is not None:
                        cells.append({"coordinate": cell.coordinate, "value_or_formula": cell.value,
                                      "cached": value.value, "type": cell.data_type})
            report["sheets"].append({"name": ws.title, "state": ws.sheet_state,
                                      "rows": ws.max_row, "columns": ws.max_column, "cells": cells})
        formulas.close()
        values.close()
        (out / (source.stem + ".json")).write_text(json.dumps(report, indent=2, default=str) + "\n")
        print(source.name, [(s["name"], s["rows"], s["columns"]) for s in report["sheets"]])


if __name__ == "__main__":
    main()
