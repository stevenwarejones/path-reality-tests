"""Strict readers for external pinned archives; never execute or extract author code."""
from dataclasses import dataclass
import hashlib
import io
import json
from pathlib import Path
import zipfile
import numpy as np
import openpyxl

HERE = Path(__file__).resolve().parent


def digest(data):
    return hashlib.sha256(data).hexdigest()


def verify_sources(directory):
    directory = Path(directory).resolve()
    if directory.is_relative_to(HERE.parents[1]): raise ValueError("Use external cache")
    manifest = json.loads((HERE/"manifest.json").read_text())
    for item in manifest["files"]:
        raw = (directory/item["name"]).read_bytes()
        if len(raw)!=item["bytes"] or digest(raw)!=item["sha256"]: raise ValueError("Source integrity failure: "+item["name"])
    return manifest


def table(raw,columns):
    a = np.loadtxt(io.BytesIO(raw),ndmin=2)
    if a.shape[1]!=columns or not np.isfinite(a).all(): raise ValueError("Invalid table")
    if not np.all(np.diff(a[:,0])>0): raise ValueError("Coordinates must increase")
    return a


@dataclass
class Scan:
    name: str
    position_m: np.ndarray
    counts: np.ndarray
    powers_w: np.ndarray
    incident_rate_hz: float
    integration_s: float


def sodium(directory):
    prefix = "Na-Cluster-Interference/Na-Cluster/Data_April2025/22042025/"
    with zipfile.ZipFile(Path(directory)/"Na-Cluster-Interference.zip") as z:
        sheet = openpyxl.load_workbook(io.BytesIO(z.read(prefix+"Powerscan_85kTh_22042025.xlsx")),data_only=True).active
        if sheet["A1"].value!="name" or sheet.max_row!=96: raise ValueError("Unexpected sodium schema")
        scans = []
        for row in sheet.iter_rows(min_row=2,values_only=True):
            name,rate,duration,steps,spacing,*rest = row
            a = table(z.read(prefix+"Interference Scans/"+name+"_s1.dat"),2)
            if len(a)!=steps or spacing<=0: raise ValueError("Invalid scan length")
            # Measured coordinates are nonuniform; 15 nm is only the nominal step.
            if np.any(a[:,1]<0) or np.any(a[:,1]!=np.floor(a[:,1])): raise ValueError("Invalid counts")
            if duration<=0 or rate<=0 or min(rest[:3])<0: raise ValueError("Invalid settings")
            scans.append(Scan(name,a[:,0]*1e-9,a[:,1],np.array(rest[:3])*1e-3,rate*1000,duration))
        mass = table(z.read(prefix+"Mass spec/Sodium_He_60sccm_Ar_161sccm_356C_425nm_L12p5cm_HP_massscan_16.dat"),2)
        if np.any(mass[:,1]<0) or mass[:,1].sum()<=0: raise ValueError("Invalid mass spectrum")
        m = mass[:,0]*2
        mi = np.linspace(m.min(),m.max(),10*len(m))
        wi = np.interp(mi,m,mass[:,1])
        keep = (mi>.85*170000)&(mi<1.15*170000)
        return scans,mi[keep],wi[keep]/wi[keep].sum()


def force(directory):
    path = Path(directory)/"Final_data.xlsx"
    sheet = openpyxl.load_workbook(path,data_only=True).active
    if [sheet.cell(1,c).value for c in range(1,6)]!=["T","Q","T/Q","B","errB"]: raise ValueError("Unexpected force schema")
    rows = np.array([[sheet.cell(row,col).value for col in range(1,6)] for row in range(3,17)],float)
    if not np.isfinite(rows).all() or np.any(rows<=0): raise ValueError("Invalid force summary")
    if not np.allclose(rows[:,0]*1e-3/rows[:,1],rows[:,2],rtol=1e-12,atol=0): raise ValueError("T/Q unit mismatch")
    spectra = {}
    with zipfile.ZipFile(Path(directory)/"PSD.zip") as z:
        for temp in rows[:,0].astype(int):
            a = table(z.read(f"PSD/{temp}.dat"),3)
            if np.any(a[:,1:]<=0): raise ValueError("Invalid PSD")
            spectra[int(temp)] = a
    cells = {c:float(sheet[c].value) for c in ["D20","E20","D23","E23","K20","L20","M20","N20","O20","L26","M26","N29"]}
    formulas = openpyxl.load_workbook(path,data_only=False).active
    return rows,spectra,cells,{c:formulas[c].value for c in ["K23","L23","M23","N23","L26","M26","N29"]}

